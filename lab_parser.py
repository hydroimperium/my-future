import io
import os
import re
import json
import hashlib
from dotenv import load_dotenv
from openai import OpenAI
from schemas import UserHabitsInput
from ai_service import search_authoritative_medical_web

load_dotenv()

# Детерминированный кэш результатов разбора анализов (ключ: SHA-256 от текста файла + возраст + пол)
_LAB_EVAL_CACHE: dict[str, dict] = {}

# Ключевые слова стандартных анализов, которые УЖЕ входят в 6 базовых кластеров
# и НЕ должны повторно попадать в unlisted_abnormal_markers (защита от двойного счёта внутри бланка)
STANDARD_CLUSTER_KEYWORDS = (
    "глюкоз", "glucose", "сахар", "гликирован", "hba1c", "a1c", "homa", "инсулин",
    "холестерин", "холестерол", "cholesterol", "лпнп", "лдл", "ldl", "лпвп", "hdl", "триглицерид", "атероген",
    "скф", "фильтрац", "egfr", "ckd-epi", "креатинин", "creatinine", "алт", "alt", "аланин",
    "аст", "ast", "аспартат", "билирубин", "bilirubin", "мочевая кислота", "uric",
    "гемоглобин", "hemoglobin", "hgb", "ферритин", "ferritin", "ттг", "tsh", "тиреотроп", "витамин d", "25-oh",
    "реактивный белок", "срб", "crp", "гомоцистеин", "homocysteine"
)


def extract_text_from_file(filename: str, file_bytes: bytes) -> str:
    """Извлекает текст и таблицы из файлов .pdf, .docx и .txt в оперативной памяти."""
    ext = filename.lower().split(".")[-1]
    extracted_lines = []

    if ext == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(file_bytes))
        for page in reader.pages[:15]:
            page_text = page.extract_text()
            if page_text:
                extracted_lines.append(page_text)

    elif ext in ("docx", "doc"):
        from docx import Document
        doc = Document(io.BytesIO(file_bytes))
        for p in doc.paragraphs:
            if p.text.strip():
                extracted_lines.append(p.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells if cell.text.strip()]
                dedup_cells = []
                for c in row_cells:
                    if not dedup_cells or dedup_cells[-1] != c:
                        dedup_cells.append(c)
                if dedup_cells:
                    extracted_lines.append(" | ".join(dedup_cells))

    else:
        text = file_bytes.decode("utf-8", errors="ignore")
        extracted_lines.append(text)

    full_text = "\n".join(extracted_lines)
    full_text = re.sub(r"[ \t]+", " ", full_text)
    return full_text[:15000]


def _strip_reference_ranges(text: str) -> str:
    """Удаляет справочные интервалы нормы и даты, чтобы регулярные выражения брали только результат пациента."""
    cleaned = re.sub(r"\b\d{2}\.\d{2}\.\d{2,4}\b", " ", text)
    cleaned = re.sub(r"\((?:норма|референс|ref)[^)]*\)", " ", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\b\d+(?:[.,]\d+)?\s*[-–—]\s*\d+(?:[.,]\d+)?\b", " ", cleaned)
    return cleaned


def _extract_biomarker_value(pattern: str, cleaned_text: str) -> float | None:
    match = re.search(pattern, cleaned_text, re.IGNORECASE)
    if not match:
        return None
    val_str = match.group(1).replace(",", ".")
    try:
        return float(val_str)
    except ValueError:
        return None


def _calculate_ckd_epi_2021(creatinine_umol: float, age: int, gender: str) -> float:
    """Формула CKD-EPI (2021) для расчёта СКФ почек (мл/мин/1.73м²)."""
    scr_mg_dl = creatinine_umol / 88.4
    if gender == "female":
        kappa, alpha, sex_mult = 0.7, -0.241, 1.012
    else:
        kappa, alpha, sex_mult = 0.9, -0.302, 1.0

    ratio = scr_mg_dl / kappa
    egfr = 142.0 * (min(ratio, 1.0) ** alpha) * (max(ratio, 1.0) ** -1.200) * (0.9938 ** age) * sex_mult
    return round(egfr, 1)


def _deterministic_biomarker_scan(raw_text: str, age: int = 45, gender: str = "male") -> dict:
    """Половозрастной детерминированный сканер базовых лабораторных показателей."""
    cleaned = _strip_reference_ranges(raw_text.lower())

    clusters = {
        "glycemic": {"impact_years": 0.0, "irreversible_fraction": 0.30, "findings": []},
        "lipids": {"impact_years": 0.0, "irreversible_fraction": 0.30, "findings": []},
        "renal_hepatic": {"impact_years": 0.0, "irreversible_fraction": 0.45, "findings": []},
        "hematology_endocrine": {"impact_years": 0.0, "irreversible_fraction": 0.20, "findings": []},
        "inflammation": {"impact_years": 0.0, "irreversible_fraction": 0.25, "findings": []},
        "coagulation_cardiac": {"impact_years": 0.0, "irreversible_fraction": 0.45, "findings": []},
    }
    normal_markers_found = []
    recommendations = []

    # 1. Углеводный обмен
    glucose = _extract_biomarker_value(r"(?:глюкоза|glucose|сахар\s+крови)[^\d\n]{1,35}?(\d{1,2}(?:[.,]\d{1,2})?)", cleaned)
    hba1c = _extract_biomarker_value(r"(?:гликирован\w+\s+гемоглобин|hba1c|a1c)[^\d\n]{1,35}?(\d{1,2}(?:[.,]\d{1,2})?)", cleaned)
    homa = _extract_biomarker_value(r"(?:homa-ir|индекс\s+инсулинорезистентности|homa)[^\d\n]{1,30}?(\d{1,2}(?:[.,]\d{1,2})?)", cleaned)

    glyc_penalties = []
    if glucose is not None and 2.2 <= glucose <= 32.0:
        if glucose < 3.3:
            glyc_penalties.append(-1.2)
            clusters["glycemic"]["findings"].append(f"Гипогликемия натощак ({glucose} ммоль/л)")
        elif glucose <= 5.5:
            normal_markers_found.append(f"Глюкоза {glucose} ммоль/л (норма)")
        elif glucose <= 6.0:
            glyc_penalties.append(-0.8)
            clusters["glycemic"]["findings"].append(f"Пограничная глюкоза ({glucose} ммоль/л)")
        elif glucose <= 6.9:
            glyc_penalties.append(-1.8)
            clusters["glycemic"]["findings"].append(f"Повышена глюкоза натощак ({glucose} ммоль/л — предиабет)")
            recommendations.append("Исключить быстрые углеводы и проверить гликированный гемоглобин (HbA1c).")
        elif glucose <= 9.0:
            glyc_penalties.append(-3.6)
            clusters["glycemic"]["irreversible_fraction"] = 0.40
            clusters["glycemic"]["findings"].append(f"Высокая глюкоза крови ({glucose} ммоль/л)")
            recommendations.append("Консультация эндокринолога для коррекции углеводного обмена.")
        else:
            glyc_penalties.append(-5.5)
            clusters["glycemic"]["irreversible_fraction"] = 0.48
            clusters["glycemic"]["findings"].append(f"Выраженная гипергликемия ({glucose} ммоль/л)")

    if hba1c is not None and 3.5 <= hba1c <= 20.0:
        if hba1c < 5.7:
            normal_markers_found.append(f"HbA1c {hba1c}% (норма)")
        elif hba1c <= 6.4:
            glyc_penalties.append(-1.8)
            clusters["glycemic"]["findings"].append(f"Повышен гликированный гемоглобин ({hba1c}% — предиабет)")
        elif hba1c <= 7.5:
            glyc_penalties.append(-3.5)
            clusters["glycemic"]["irreversible_fraction"] = 0.40
            clusters["glycemic"]["findings"].append(f"HbA1c {hba1c}% (диабетический диапазон)")
        else:
            glyc_penalties.append(-5.5)
            clusters["glycemic"]["irreversible_fraction"] = 0.48
            clusters["glycemic"]["findings"].append(f"Декомпенсированный HbA1c ({hba1c}%)")

    if homa is not None and 0.3 <= homa <= 25.0:
        if homa > 2.7:
            glyc_penalties.append(-1.2)
            clusters["glycemic"]["findings"].append(f"Повышен индекс HOMA-IR ({homa})")

    if glyc_penalties:
        clusters["glycemic"]["impact_years"] = round(min(glyc_penalties), 2)

    # 2. Липидный спектр
    chol = _extract_biomarker_value(r"(?:холестерин\s+общ\w*|холестерол\s+общ\w*|total\s+cholesterol|холестерин)[^\d\n]{1,35}?(\d{1,2}[.,]\d{1,2})", cleaned)
    ldl = _extract_biomarker_value(r"(?:лпнп|лдл|ldl|низкой\s+плотности)[^\d\n]{1,35}?(\d{1,2}[.,]\d{1,2})", cleaned)
    hdl = _extract_biomarker_value(r"(?:лпвп|hdl|высокой\s+плотности)[^\d\n]{1,35}?(\d{1,2}[.,]\d{1,2})", cleaned)
    trig = _extract_biomarker_value(r"(?:триглицерид\w*|triglycerides)[^\d\n]{1,35}?(\d{1,2}[.,]\d{1,2})", cleaned)

    lipid_penalties = []
    if chol is not None and 2.0 <= chol <= 18.0:
        if chol <= 5.1:
            normal_markers_found.append(f"Общий холестерин {chol} ммоль/л (норма)")
        elif chol <= 6.2:
            lipid_penalties.append(-1.1)
            clusters["lipids"]["findings"].append(f"Умеренно повышен общий холестерин ({chol} ммоль/л)")
        elif chol <= 7.5:
            lipid_penalties.append(-2.2)
            clusters["lipids"]["findings"].append(f"Повышенный общий холестерин ({chol} ммоль/л)")
            recommendations.append("Скорректировать жировой рацион и контролировать липидограмму.")
        else:
            lipid_penalties.append(-3.4)
            clusters["lipids"]["irreversible_fraction"] = 0.35
            clusters["lipids"]["findings"].append(f"Выраженная гиперхолестеринемия ({chol} ммоль/л)")

    if ldl is not None and 0.8 <= ldl <= 12.0:
        if ldl <= 3.0:
            normal_markers_found.append(f"ЛПНП {ldl} ммоль/л (норма)")
        elif ldl <= 4.0:
            lipid_penalties.append(-1.3)
            clusters["lipids"]["findings"].append(f"Повышен холестерин ЛПНП ({ldl} ммоль/л)")
        elif ldl <= 5.0:
            lipid_penalties.append(-2.4)
            clusters["lipids"]["findings"].append(f"Высокий уровень ЛПНП ({ldl} ммоль/л)")
        else:
            lipid_penalties.append(-3.6)
            clusters["lipids"]["findings"].append(f"Критически высокий ЛПНП ({ldl} ммоль/л)")

    if hdl is not None and 0.3 <= hdl <= 4.0:
        hdl_min = 1.0 if gender == "male" else 1.2
        if hdl < hdl_min:
            lipid_penalties.append(-1.0)
            clusters["lipids"]["findings"].append(f"Снижен защитный холестерин ЛПВП ({hdl} ммоль/л)")

    if trig is not None and 0.3 <= trig <= 15.0:
        if trig > 2.2:
            lipid_penalties.append(-1.3 if trig <= 4.0 else -2.2)
            clusters["lipids"]["findings"].append(f"Повышены триглицериды ({trig} ммоль/л)")

    if lipid_penalties:
        clusters["lipids"]["impact_years"] = round(max(-3.8, min(lipid_penalties)), 2)

    # 3. Почки и печень
    egfr = _extract_biomarker_value(r"(?:скф|клубочков\w+\s+фильтрац\w*|egfr|ckd-epi)[^\d\n]{1,35}?(\d{2,3}(?:[.,]\d)?)", cleaned)
    creat = _extract_biomarker_value(r"(?:креатинин|creatinine)[^\d\n]{1,35}?(\d{2,3}(?:[.,]\d)?)", cleaned)
    alt = _extract_biomarker_value(r"(?:\bалт\b|аланинаминотрансфераза|\balt\b)[^\d\n]{1,35}?(\d{1,3}(?:[.,]\d)?)", cleaned)
    ast = _extract_biomarker_value(r"(?:\bаст\b|аспартатаминотрансфераза|\bast\b)[^\d\n]{1,35}?(\d{1,3}(?:[.,]\d)?)", cleaned)
    bili = _extract_biomarker_value(r"(?:билирубин\s+общ\w*|билирубин|bilirubin)[^\d\n]{1,35}?(\d{1,3}[.,]\d{1,2})", cleaned)
    uric = _extract_biomarker_value(r"(?:мочевая\s+кислота|uric\s+acid)[^\d\n]{1,35}?(\d{2,3}(?:[.,]\d)?)", cleaned)

    rh_penalties = []
    if egfr is None and creat is not None and 35 <= creat <= 900:
        egfr = _calculate_ckd_epi_2021(creat, age, gender)

    if egfr is not None and 5 <= egfr <= 160:
        if egfr >= 75:
            normal_markers_found.append(f"Расчётная СКФ {egfr} мл/мин (норма)")
        elif egfr >= 60:
            rh_penalties.append(-0.7)
            clusters["renal_hepatic"]["findings"].append(f"Умеренное снижение СКФ ({egfr} мл/мин)")
        elif egfr >= 45:
            rh_penalties.append(-2.5)
            clusters["renal_hepatic"]["irreversible_fraction"] = 0.55
            clusters["renal_hepatic"]["findings"].append(f"Снижение фильтрационной функции почек (СКФ {egfr} мл/мин)")
            recommendations.append("Контроль функции почек у врача, ограничение соли и отказ от НПВС.")
        else:
            rh_penalties.append(-5.8)
            clusters["renal_hepatic"]["irreversible_fraction"] = 0.70
            clusters["renal_hepatic"]["findings"].append(f"Выраженное снижение функции почек (СКФ {egfr} мл/мин)")

    max_trans = max(alt or 0.0, ast or 0.0)
    if 8 <= max_trans <= 1500:
        if max_trans <= 40:
            normal_markers_found.append(f"АЛТ/АСТ в норме ({max_trans} Ед/л)")
        elif max_trans <= 85:
            rh_penalties.append(-1.1)
            clusters["renal_hepatic"]["findings"].append(f"Повышены печеночные ферменты АЛТ/АСТ (до {max_trans} Ед/л)")
            recommendations.append("Выполнить УЗИ печени и снизить метаболическую/алкогольную нагрузку.")
        else:
            rh_penalties.append(-2.5)
            clusters["renal_hepatic"]["findings"].append(f"Выраженное повышение АЛТ/АСТ ({max_trans} Ед/л)")

    if bili is not None and 2.0 <= bili <= 300.0:
        if bili > 22.0:
            rh_penalties.append(-0.9 if bili <= 40.0 else -2.2)
            clusters["renal_hepatic"]["findings"].append(f"Повышен общий билирубин ({bili} мкмоль/л)")

    if uric is not None and 100 <= uric <= 900:
        uric_limit = 420.0 if gender == "male" else 360.0
        if uric > uric_limit:
            rh_penalties.append(-1.0 if uric <= uric_limit + 100 else -1.8)
            clusters["renal_hepatic"]["findings"].append(f"Повышена мочевая кислота ({uric} мкмоль/л)")

    if rh_penalties:
        clusters["renal_hepatic"]["impact_years"] = round(max(-7.5, min(rh_penalties)), 2)

    # 4. Кроветворение, дефициты и щитовидная железа
    hb = _extract_biomarker_value(r"(?:гемоглобин(?!\s+глик)|hemoglobin|\bhgb\b|\bhb\b)[^\d\n]{1,35}?(\d{2,3}(?:[.,]\d)?)", cleaned)
    ferritin = _extract_biomarker_value(r"(?:ферритин|ferritin)[^\d\n]{1,35}?(\d{1,4}(?:[.,]\d)?)", cleaned)
    tsh = _extract_biomarker_value(r"(?:\bттг\b|тиреотропн\w+|\btsh\b)[^\d\n]{1,35}?(\d{1,2}[.,]\d{1,2})", cleaned)
    vit_d = _extract_biomarker_value(r"(?:витамин\s*d|25-oh|25\(oh\)d)[^\d\n]{1,35}?(\d{1,3}(?:[.,]\d)?)", cleaned)

    hem_penalties = []
    if hb is not None and 50 <= hb <= 220:
        hb_low_norm = 130.0 if gender == "male" else 120.0
        hb_high_norm = 168.0 if gender == "male" else 155.0
        if hb_low_norm <= hb <= hb_high_norm:
            normal_markers_found.append(f"Гемоглобин {hb} г/л (норма)")
        elif hb_low_norm - 20 <= hb < hb_low_norm:
            hem_penalties.append(-1.1)
            clusters["hematology_endocrine"]["findings"].append(f"Снижен гемоглобин ({hb} г/л — легкая анемия)")
            recommendations.append("Проверить ферритин, сывороточное железо и витамин B12.")
        elif hb < hb_low_norm - 20:
            hem_penalties.append(-2.3)
            clusters["hematology_endocrine"]["findings"].append(f"Выраженная анемия (гемоглобин {hb} г/л)")
        elif hb > hb_high_norm:
            hem_penalties.append(-1.4)
            clusters["hematology_endocrine"]["findings"].append(f"Повышен гемоглобин ({hb} г/л — сгущение крови)")
            recommendations.append("Соблюдать питьевой режим и контролировать вязкость крови.")

    if ferritin is not None and 1 <= ferritin <= 1500:
        ferr_min = 30.0 if gender == "male" else 18.0
        if ferritin < ferr_min:
            hem_penalties.append(-0.6)
            clusters["hematology_endocrine"]["findings"].append(f"Снижен ферритин ({ferritin} мкг/л — дефицит запасов железа)")
        elif ferr_min <= ferritin <= 300:
            normal_markers_found.append(f"Ферритин {ferritin} мкг/л (норма)")

    if tsh is not None and 0.01 <= tsh <= 50.0:
        if 0.4 <= tsh <= 4.2:
            normal_markers_found.append(f"ТТГ {tsh} мМЕ/л (норма)")
        elif tsh > 4.2:
            hem_penalties.append(-1.0 if tsh <= 8.0 else -1.6)
            clusters["hematology_endocrine"]["findings"].append(f"Повышен ТТГ ({tsh} мМЕ/л — гипофункция щитовидной железы)")
        else:
            hem_penalties.append(-1.2)
            clusters["hematology_endocrine"]["findings"].append(f"Снижен ТТГ ({tsh} мМЕ/л)")

    if vit_d is not None and 3 <= vit_d <= 150:
        if vit_d < 20:
            hem_penalties.append(-0.5)
            clusters["hematology_endocrine"]["findings"].append(f"Дефицит витамина D ({vit_d} нг/мл)")

    if hem_penalties:
        clusters["hematology_endocrine"]["impact_years"] = round(max(-2.6, min(hem_penalties)), 2)

    # 5. Воспаление
    crp = _extract_biomarker_value(r"(?:с-реактивн\w+\s+белок|\bсрб\b|\bcrp\b)[^\d\n]{1,35}?(\d{1,3}(?:[.,]\d{1,2})?)", cleaned)
    homocyst = _extract_biomarker_value(r"(?:гомоцистеин|homocysteine)[^\d\n]{1,35}?(\d{1,2}(?:[.,]\d{1,2})?)", cleaned)

    inf_penalties = []
    if crp is not None and 0.1 <= crp <= 200.0:
        if crp <= 2.5:
            normal_markers_found.append(f"С-реактивный белок {crp} мг/л (норма)")
        elif crp <= 8.0:
            inf_penalties.append(-1.2)
            clusters["inflammation"]["findings"].append(f"Повышен С-реактивный белок ({crp} мг/л)")
        else:
            inf_penalties.append(-2.2)
            clusters["inflammation"]["findings"].append(f"Высокий С-реактивный белок ({crp} мг/л)")

    if homocyst is not None and 3.0 <= homocyst <= 60.0:
        if homocyst > 13.5:
            inf_penalties.append(-1.3)
            clusters["inflammation"]["findings"].append(f"Повышен гомоцистеин ({homocyst} мкмоль/л)")

    if inf_penalties:
        clusters["inflammation"]["impact_years"] = round(max(-2.4, min(inf_penalties)), 2)

    return {
        "clusters": clusters,
        "normal_markers_found": normal_markers_found,
        "recommendations": recommendations
    }


def evaluate_lab_report_for_prediction(lab_upload_data: dict, user_data: UserHabitsInput) -> dict | None:
    """
    Выполняет клинический расчёт прикрепленного файла анализов с SHA-256 кэшированием
    (чтобы от итерации к итерации цифры одного и того же файла не менялись).
    """
    if not lab_upload_data or lab_upload_data.get("status") != "ok":
        return None

    filename = lab_upload_data.get("filename", "Анализы")
    raw_text = lab_upload_data.get("raw_text", "")
    if not raw_text or len(raw_text.strip()) < 5:
        return None

    # Проверяем детерминированный кэш
    cache_key = hashlib.sha256(f"{raw_text}|{user_data.age}|{user_data.gender}".encode("utf-8")).hexdigest()
    if cache_key in _LAB_EVAL_CACHE:
        return copy_lab_dict(_LAB_EVAL_CACHE[cache_key])

    base_scan = _deterministic_biomarker_scan(raw_text, age=user_data.age, gender=user_data.gender)
    clusters = base_scan["clusters"]
    normal_markers = base_scan["normal_markers_found"]
    recommendations = base_scan["recommendations"]

    api_key = os.getenv("LLM_API_KEY", "").strip()
    base_url = os.getenv("LLM_BASE_URL", "https://generativelanguage.googleapis.com/v1beta/openai/").strip()
    primary_model = os.getenv("LLM_MODEL", "gemini-3.5-flash-lite").strip()

    llm_summary = ""
    llm_action = ""
    web_searched_markers_sources = []
    unlisted_impact_years = 0.0
    unlisted_irrev_fraction = 0.35

    if api_key and api_key != "ПУСТО":
        client = OpenAI(api_key=api_key, base_url=base_url, timeout=18.0)
        candidate_models = [primary_model]
        for backup_m in ["gemini-3.5-flash-lite", "gemini-3.8-flash"]:
            if backup_m not in candidate_models:
                candidate_models.append(backup_m)

        extract_prompt = f"""
Ты — врач лабораторной диагностики. Пациент: возраст {user_data.age} лет, пол {user_data.gender}.
Проанализируй текст бланка анализов:
---
{raw_text[:8000]}
---

ЗАДАЧА:
1. Распредели найденные отклонения по 6 базовым кластерам (`glycemic`, `lipids`, `renal_hepatic`, `hematology_endocrine`, `inflammation`, `coagulation_cardiac`).
2. В список `unlisted_abnormal_markers` запиши ТОЛЬКО те показатели с отклонением от нормы, которых НЕТ среди стандартных (НЕ записывай туда глюкозу, сахар, HbA1c, холестерин, ЛПНП, ЛПВП, триглицериды, креатинин, СКФ, АЛТ, АСТ, билирубин, мочевую кислоту, гемоглобин, ферритин, ТТГ, витамин D, СРБ, гомоцистеин).

Ответь СТРОГО валидным JSON:
{{
  "clusters": {{
    "glycemic": {{"impact_years": -1.8, "findings": ["Глюкоза 6.4 ммоль/л"]}},
    "lipids": {{"impact_years": 0.0, "findings": []}},
    "renal_hepatic": {{"impact_years": 0.0, "findings": []}},
    "hematology_endocrine": {{"impact_years": 0.0, "findings": []}},
    "inflammation": {{"impact_years": 0.0, "findings": []}},
    "coagulation_cardiac": {{"impact_years": 0.0, "findings": []}}
  }},
  "unlisted_abnormal_markers": [],
  "normal_markers": ["Гемоглобин 145 г/л"],
  "clinical_summary": "Краткое резюме по анализам на 2 предложения простым языком.",
  "recommended_action": "Конкретная рекомендация (1 предложение)."
}}
"""
        unlisted_markers = []
        for model_to_try in candidate_models:
            try:
                resp = client.chat.completions.create(
                    model=model_to_try,
                    messages=[
                        {"role": "system", "content": "Ты — врач лабораторной диагностики. Отвечай строго чистым JSON."},
                        {"role": "user", "content": extract_prompt}
                    ],
                    temperature=0.0
                )
                raw_reply = resp.choices[0].message.content.strip()
                json_match = re.search(r"\{.*\}", raw_reply, re.DOTALL)
                if json_match:
                    parsed = json.loads(json_match.group(0))
                    llm_clusters = parsed.get("clusters", {})
                    cluster_caps = {
                        "glycemic": -5.8, "lipids": -3.8, "renal_hepatic": -7.5,
                        "hematology_endocrine": -2.6, "inflammation": -2.4, "coagulation_cardiac": -4.2
                    }
                    for c_key, cap in cluster_caps.items():
                        if c_key in llm_clusters:
                            llm_val = float(llm_clusters[c_key].get("impact_years", 0.0))
                            det_val = clusters[c_key]["impact_years"]
                            # Если наш детерминированный сканер уже вычислил точный сдвиг по формуле, берем его как основной якорь
                            if det_val < 0:
                                merged_val = det_val
                            else:
                                merged_val = llm_val * 0.85
                            clusters[c_key]["impact_years"] = round(max(cap, min(0.0, merged_val)), 2)
                            llm_findings = llm_clusters[c_key].get("findings", [])
                            if llm_findings and not clusters[c_key]["findings"]:
                                clusters[c_key]["findings"] = [str(x) for x in llm_findings[:3]]

                    # Фильтруем unlisted_abnormal_markers от случайного попадания стандартных анализов
                    raw_unlisted = [str(m) for m in parsed.get("unlisted_abnormal_markers", []) if str(m).strip()]
                    unlisted_markers = [
                        m for m in raw_unlisted
                        if not any(std_kw in m.lower() for std_kw in STANDARD_CLUSTER_KEYWORDS)
                    ]

                    for nm in parsed.get("normal_markers", []):
                        if nm not in normal_markers and len(normal_markers) < 8:
                            normal_markers.append(str(nm))

                    llm_summary = str(parsed.get("clinical_summary", "")).strip()
                    llm_action = str(parsed.get("recommended_action", "")).strip()
                    break
            except Exception as e:
                print(f"[Предупреждение LLM Lab Extract ({model_to_try})]: {e}")

        if unlisted_markers:
            search_q = "; ".join(unlisted_markers[:3])
            web_context, web_sources, _ = search_authoritative_medical_web(search_q)
            web_searched_markers_sources = web_sources

            web_eval_prompt = f"""
Ты — специалист по доказательной медицине. В бланке анализов пациента ({user_data.age} лет, пол {user_data.gender}) обнаружены редкие/дополнительные показатели вне базового шаблона:
{search_q}

АВТОРИТЕТНЫЕ ДАННЫЕ ИЗ МЕДИЦИНСКОГО ПОИСКА (ВОЗ, PubMed, Клинические рекомендации):
{web_context}

Оцени суммарный чистый вклад этих показателей в прогноз продолжительности жизни (`unlisted_impact_years` от -4.5 до 0.0 лет) и долю необратимости (`irreversible_fraction` от 0.15 до 0.75).
Ответь СТРОГО чистым JSON:
{{
  "unlisted_impact_years": -1.4,
  "irreversible_fraction": 0.35,
  "unlisted_explanation": "Краткое объяснение на 1 предложение по результатам поиска.",
  "unlisted_action": "Рекомендация по контролю."
}}
"""
            for model_to_try in candidate_models:
                try:
                    resp2 = client.chat.completions.create(
                        model=model_to_try,
                        messages=[
                            {"role": "system", "content": "Опирайся строго на авторитетные медицинские источники. Отвечай чистым JSON."},
                            {"role": "user", "content": web_eval_prompt}
                        ],
                        temperature=0.0
                    )
                    m2 = re.search(r"\{.*\}", resp2.choices[0].message.content.strip(), re.DOTALL)
                    if m2:
                        p2 = json.loads(m2.group(0))
                        unlisted_impact_years = round(max(-4.5, min(0.0, float(p2.get("unlisted_impact_years", -1.0)))), 2)
                        unlisted_irrev_fraction = round(max(0.15, min(0.80, float(p2.get("irreversible_fraction", 0.35)))), 2)
                        extra_exp = str(p2.get("unlisted_explanation", "")).strip()
                        if extra_exp:
                            llm_summary = f"{llm_summary} {extra_exp}".strip()
                        break
                except Exception as e:
                    print(f"[Предупреждение Web Eval Unlisted Labs ({model_to_try})]: {e}")

        for um in unlisted_markers:
            if um not in clusters["coagulation_cardiac"]["findings"]:
                clusters["coagulation_cardiac"]["findings"].append(f"{um} (проверено по базе мед. источников)")

    detected_abnormalities = []
    for c_key, c_data in clusters.items():
        for f_str in c_data.get("findings", []):
            if f_str not in detected_abnormalities:
                detected_abnormalities.append(f_str)

    all_normal_bonus = 0.0
    if not detected_abnormalities and len(normal_markers) >= 2:
        all_normal_bonus = 0.8

    if not llm_summary:
        if detected_abnormalities:
            llm_summary = f"По результатам разбора файла выявлены отклонения: {'; '.join(detected_abnormalities)}."
        elif normal_markers:
            llm_summary = f"Все проверенные лабораторные показатели находятся в пределах половозрастной нормы ({', '.join(normal_markers[:4])})."
        else:
            llm_summary = "Файл анализов обработан."

    if not llm_action:
        llm_action = recommendations[0] if recommendations else "Обсудить результаты лабораторной панели с лечащим врачом."

    result = {
        "filename": filename,
        "status": "ok",
        "clusters": clusters,
        "unlisted_impact_years": unlisted_impact_years,
        "unlisted_irrev_fraction": unlisted_irrev_fraction,
        "web_sources": web_searched_markers_sources,
        "detected_abnormalities": detected_abnormalities,
        "normal_markers": normal_markers,
        "all_normal_bonus_years": all_normal_bonus,
        "clinical_summary": llm_summary,
        "recommended_action": llm_action
    }
    _LAB_EVAL_CACHE[cache_key] = copy_lab_dict(result)
    return result


def copy_lab_dict(d: dict) -> dict:
    return json.loads(json.dumps(d))
